# -*- coding: utf-8 -*-
"""生成 γ=1 可交付储能计划（同组场景集）的 result2.xlsx。

依赖 compare_deliverable.py（位于上级目录 retry/），
计划层用 solve_deliver(γ=1)，执行层用带 SOC 跟踪的诚实执行，
按附件5模板填写计划购电量 / 充放电量 / 紧急购电量。
"""
import os
import sys
import datetime as dt

import numpy as np
import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import compare_deliverable as CD


def full_honest_sim(g, c, d):
    """诚实执行并返回完整轨迹（c_act, d_act, e_act, soc_traj）。"""
    soc = CD.SOC0
    c_act = np.zeros(CD.T); d_act = np.zeros(CD.T); e_act = np.zeros(CD.T)
    soc_traj = np.zeros(CD.T + 1); soc_traj[0] = CD.SOC0
    for t in range(CD.T):
        dmax = max(0.0, (soc - CD.SOC_MIN) * CD.ETA / CD.DT)
        dt_ = min(d[t], dmax)
        cmax_soc = max(0.0, (CD.SOC_MAX - soc) / (CD.ETA * CD.DT))
        ct_ = min(c[t], max(0.0, CD.P[t] + g[t] + dt_ - CD.L[t]), cmax_soc)
        e_act[t] = max(0.0, CD.L[t] - g[t] - CD.P[t] - dt_)
        c_act[t] = ct_; d_act[t] = dt_
        soc += (CD.ETA * ct_ - dt_ / CD.ETA) * CD.DT
        soc_traj[t + 1] = soc
    return c_act, d_act, e_act, soc_traj


def fmt_time(m):
    m = int(m)
    return "24:00" if m >= 1440 else f"{m // 60:02d}:{m % 60:02d}"


def segments(e_kwh):
    segs = []; i = 0
    while i < CD.N:
        if e_kwh[i] > 1e-6:
            j = i
            while j + 1 < CD.N and e_kwh[j + 1] > 1e-6:
                j += 1
            segs.append((fmt_time(i * 10) + "-" + fmt_time((j + 1) * 10),
                         float(e_kwh[i:j + 1].sum())))
            i = j + 1
        else:
            i += 1
    return segs


if __name__ == "__main__":
    # γ=1 可交付计划（同组场景集）
    peers = CD.build_peers_group()
    g, d_plan, cm, soc_plan, el = CD.solve_deliver(peers, 1.0)
    c_act, d_act, e_act, soc_traj = full_honest_sim(g, cm, d_plan)

    N, T, NDAYS, REPORT = CD.N, CD.T, CD.NDAYS, CD.REPORT
    g_day = g.reshape(NDAYS, N) * CD.DT
    c_day = c_act.reshape(NDAYS, N) * CD.DT
    d_day = d_act.reshape(NDAYS, N) * CD.DT
    e_day = e_act.reshape(NDAYS, N) * CD.DT
    price_day = CD.price_day

    out_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    template = os.path.join(CD.DATA, "附件5", "result2.xlsx")
    out_path = os.path.join(out_dir, "结果", "result2_deliverable_group.xlsx")
    wb = openpyxl.load_workbook(template)
    report_days = range(REPORT, NDAYS)
    blocks = [("0:00-4:00", 0, 24), ("4:00-8:00", 24, 48), ("8:00-12:00", 48, 72),
              ("12:00-16:00", 72, 96), ("16:00-20:00", 96, 120), ("20:00-24:00", 120, 144)]

    ws_p = wb["计划购电量"]
    for j, di in enumerate(report_days):
        row = j + 2
        for k in range(N):
            ws_p.cell(row=row, column=2 + k).value = round(float(g_day[di][(k + 1) % N]), 4)
        ws_p.cell(row=row, column=2 + N).value = round(float(g_day[di].sum()), 4)
        ws_p.cell(row=row, column=3 + N).value = round(float((price_day * g_day[di]).sum()
                                                            + (5 * price_day * e_day[di]).sum()), 2)

    ws_c = wb["充放电量"]
    ws_c.delete_rows(2, ws_c.max_row - 1)
    date0 = dt.date(2025, 1, 1)
    soc_mid = soc_traj[::N]
    for j, di in enumerate(report_days):
        base = j * 6 + 2
        dv = date0 + dt.timedelta(days=di)
        for bj, (name, a, b) in enumerate(blocks):
            r = base + bj
            ws_c.cell(row=r, column=1).value = dv if bj == 0 else None
            ws_c.cell(row=r, column=2).value = name
            ws_c.cell(row=r, column=3).value = round(float(c_day[di][a:b].sum()), 4)
            ws_c.cell(row=r, column=4).value = round(float(d_day[di][a:b].sum()), 4)
        ws_c.cell(row=base, column=5).value = "0:00"
        ws_c.cell(row=base, column=6).value = round(float(soc_mid[di]), 4)
        ws_c.cell(row=base + 1, column=5).value = "24:00"
        ws_c.cell(row=base + 1, column=6).value = round(float(soc_mid[di + 1]), 4)

    ws_e = wb["紧急购电量"]
    ws_e.delete_rows(2, ws_e.max_row - 1)
    r = 2
    for j, di in enumerate(report_days):
        segs = segments(e_day[di])
        ws_e.cell(row=r, column=1).value = date0 + dt.timedelta(days=di)
        if segs:
            for tstr, kwh in segs:
                ws_e.cell(row=r, column=2).value = tstr
                ws_e.cell(row=r, column=3).value = round(kwh, 4)
                r += 1
        else:
            ws_e.cell(row=r, column=3).value = 0.0
            r += 1

    wb.save(out_path)
    print(f"已写入: {out_path}")
