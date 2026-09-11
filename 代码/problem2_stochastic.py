# -*- coding: utf-8 -*-
"""
问题 2 —— 无预报下的"信息论最优"策略：两阶段随机规划（Two-Stage Stochastic LP）。

动机（把"试出来的最优"升级成"可证的最优"）：
  "先点预测 → 解 LP → 滚动实际算紧急购电"这个两步流程，在 5 倍紧急购电的不对称损失下
  数学上不是最优的。因为目标对"买少了"(5 倍) 和 "买多了"(1 倍) 不对称，而均值点预测只
  对"平均误差最小"负责、不对"总费用最小"负责。

正确的无预报策略：0:00 只有历史（没有附件 3 天气预报），则用历史形成的"条件经验分布"
  直接最小化【期望】总费用。即两阶段随机规划：

    第一阶段(here-and-now，一天只取一份，不依赖具体场景)：
        计划购电 ĝ、充电 ĉ、放电 d̂、储能 SOC 递推；
    第二阶段(recourse，每个场景 ω 各一份)：
        实际负荷/光伏偏离，缺口由 5 倍价紧急购电 e_{t,ω} 补足。

    场景集 Ω_d = 与当天"同星期几"的历史天（与'同星期几'预测用同一套信息，绝不用未来天）。

  为什么可证最优：
    1) 决策变量与"点预测方案"完全相同，只是目标从"在均值点最小费用"改成"对经验分布
       最小期望费用"。点预测方案的 (ĝ*,ĉ*,d̂*) 是它的可行解 ⇒ 随机规划的期望费用 ≤
       点预测方案的期望费用（严格不差）。
    2) 它仍是线性规划 ⇒ HiGHS 报 Optimal 即全局最优。
    3) 它给的是"给定无预报信息集下的最优"，而非"16 个候选里最好的"。

输出：完美预见下界 / 点预测基线(复现16.3M) / 随机规划(期望费用 + 实际轨迹费用) 三方对比。
"""
import os
import time
import datetime as dt
import numpy as np
import pandas as pd
from scipy.optimize import linprog
from scipy import sparse
import openpyxl

DT = 1.0 / 6.0
ETA = 0.9
P_MAX = 5000.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC0 = 6000.0
N = 144
NDAYS = 365
T = N * NDAYS
REPORT = 31                    # 统计窗口 2/1–12/31（1 月暖机）
K_PEERS = 4                    # 同星期几取最近 K 周（与点预测"同星期几4周"同一信息集）

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---------------- 读数据 ----------------
df1 = pd.read_excel(os.path.join(BASE, "附件", "附件1.xlsx"))
price_day = df1.iloc[:, 1].to_numpy(float)
dfL = pd.read_excel(os.path.join(BASE, "附件", "附件2.xlsx"), sheet_name="小区负载")
dfP = pd.read_excel(os.path.join(BASE, "附件", "附件2.xlsx"), sheet_name="光伏发电实际功率")
load = dfL.iloc[:, 1:1 + N].to_numpy(float)          # (365,144) 实际负荷
pv = dfP.iloc[:, 1:1 + N].to_numpy(float)            # (365,144) 实际光伏

price = np.tile(price_day, NDAYS)
mean_load, mean_pv = load.mean(axis=0), pv.mean(axis=0)
win = np.zeros(T, dtype=bool); win[REPORT * N:] = True

# ---------------- 场景集（同星期几最近 K 周，只用 0:00 之前的信息） ----------------
peers = []
for d in range(NDAYS):
    ps = [d - 7 * k for k in range(1, K_PEERS + 1) if d - 7 * k >= 0]
    peers.append(ps)

# ---------------- 通用：全年连续 LP（计划阶段），b_eq 传"净负荷预测" N̂ = L̂ − P̂ ----------------
def solve_full_year(b_eq_net):
    """确定性 LP：min Σ price·g，s.t. 功率平衡 g+P̂+d=L̂+c+s，SOC 递推。返回 g,c,d。"""
    G0, C0, D0, S0, S_idx = 0, T, 2 * T, 3 * T, 4 * T
    n_vars = 5 * T + 1
    c_obj = np.zeros(n_vars); c_obj[G0:G0 + T] = price
    t = np.arange(T)
    rows = np.concatenate([t, t, t, t, T + t, T + t, T + t, T + t])
    cols = np.concatenate([G0 + t, C0 + t, D0 + t, S0 + t,
                           S_idx + t + 1, S_idx + t, C0 + t, D0 + t])
    data = np.concatenate([np.ones(T), -np.ones(T), np.ones(T), -np.ones(T),
                           np.ones(T), -np.ones(T), -ETA * DT * np.ones(T), (DT / ETA) * np.ones(T)])
    A_eq = sparse.coo_matrix((data, (rows, cols)), shape=(2 * T, n_vars)).tocsr()
    b_eq = np.concatenate([b_eq_net, np.zeros(T)])
    bounds = ([(0, None)] * T + [(0, P_MAX)] * T + [(0, P_MAX)] * T + [(0, None)] * T
              + [(SOC0, SOC0)] + [(SOC_MIN, SOC_MAX)] * (T - 1) + [(SOC0, SOC0)])
    res = linprog(c_obj, A_eq=A_eq, b_eq=b_eq, bounds=bounds, method="highs")
    assert res.success, res.message
    x = res.x
    return x[G0:G0 + T], x[C0:C0 + T], x[D0:D0 + T], res.fun


def rollout_cost(g, c, d):
    """把计划 (g,ĉ,d̂) 作用在实际 2025 轨迹上：紧急购电 e=max(0, L+ĉ−ĝ−P−d̂)，算报告窗口总费。"""
    e = np.maximum(0.0, load.ravel() + c - g - pv.ravel() - d)
    planned = (price * g * DT)[win].sum()
    emerg = (5 * price * e * DT)[win].sum()
    return planned + emerg, (e * DT)[win].sum()


# ================= 1) 完美预见下界（全年确定性 LP，用实际负荷/光伏） =================
g_pf, c_pf, d_pf, fun_pf = solve_full_year(load.ravel() - pv.ravel())
cost_pf, em_pf = rollout_cost(g_pf, c_pf, d_pf)

# ================= 2) 点预测基线（负荷同星期几4周 + 光伏近7天，复现 16.3M） =================
L_hat = np.empty_like(load)
for d in range(NDAYS):
    idx = [d - 7 * k for k in range(1, 5) if d - 7 * k >= 0]
    L_hat[d] = load[idx].mean(axis=0) if idx else mean_load
P_hat = np.empty_like(pv)
for d in range(NDAYS):
    lo = max(0, d - 7)
    P_hat[d] = pv[lo:d].mean(axis=0) if lo < d else mean_pv
g_pt, c_pt, d_pt, fun_pt = solve_full_year(L_hat.ravel() - P_hat.ravel())
cost_pt, em_pt = rollout_cost(g_pt, c_pt, d_pt)

# ================= 3) 两阶段随机规划（无预报最优） =================
t0 = time.time()
G0, C0, D0, S_idx = 0, T, 2 * T, 3 * T
E0 = 4 * T + 1                          # SOC 占 [3T, 4T]，共 T+1 个；紧急购电从 4T+1 起
E_total = sum(len(ps) * N for ps in peers)
n_vars = E0 + E_total
print(f"[随机规划] 场景总数 {E_total:,}，变量 {n_vars:,}，约束 {T + E_total:,}")

base_e = np.zeros(NDAYS, dtype=np.int64)
cnt = 0
for d in range(NDAYS):
    base_e[d] = cnt
    cnt += len(peers[d]) * N
assert cnt == E_total

c_obj = np.zeros(n_vars)
c_obj[G0:G0 + T] = price                # 计划购电费
# 紧急购电目标系数 = 5·price / K_d（对场景求期望）
for d in range(NDAYS):
    Kd = len(peers[d])
    if Kd == 0:
        continue
    for w in range(Kd):
        sl = slice(base_e[d] + w * N, base_e[d] + (w + 1) * N)
        c_obj[E0 + sl.start:E0 + sl.stop] = 5.0 * price_day / Kd

# --- 等式：SOC 递推（T 行，每行 4 个非零） ---
t = np.arange(T)
eq_rows = np.tile(t, 4)
eq_cols = np.concatenate([S_idx + t + 1, S_idx + t, C0 + t, D0 + t])
eq_data = np.concatenate([np.ones(T), -np.ones(T), -ETA * DT * np.ones(T), (DT / ETA) * np.ones(T)])
A_eq = sparse.coo_matrix((eq_data, (eq_rows, eq_cols)), shape=(T, n_vars)).tocsr()
b_eq = np.zeros(T)

# --- 不等式：每个场景紧急购电 e_{t,ω} ≥ L+ĉ−ĝ−P−d̂，即 −g −d +c −e ≤ P−L（E_total 行） ---
nz = 4 * E_total
ub_row = np.empty(nz, dtype=np.int64)
ub_col = np.empty(nz, dtype=np.int64)
ub_dat = np.empty(nz, dtype=np.float64)
ub_b = np.empty(E_total, dtype=np.float64)
p = 0
for d in range(NDAYS):
    Kd = len(peers[d])
    for w, pd_ in enumerate(peers[d]):
        sl = slice(base_e[d] + w * N, base_e[d] + (w + 1) * N)
        r = np.arange(sl.start, sl.stop)
        t = d * N + np.arange(N)
        ub_row[p:p + N] = r; ub_col[p:p + N] = G0 + t; ub_dat[p:p + N] = -1.0; p += N
        ub_row[p:p + N] = r; ub_col[p:p + N] = D0 + t; ub_dat[p:p + N] = -1.0; p += N
        ub_row[p:p + N] = r; ub_col[p:p + N] = C0 + t; ub_dat[p:p + N] = +1.0; p += N
        ub_row[p:p + N] = r; ub_col[p:p + N] = E0 + r;  ub_dat[p:p + N] = -1.0; p += N
        ub_b[sl] = pv[pd_] - load[pd_]
A_ub = sparse.coo_matrix((ub_dat, (ub_row, ub_col)), shape=(E_total, n_vars)).tocsr()

bounds = ([(0, None)] * T + [(0, P_MAX)] * T + [(0, P_MAX)] * T          # g, c, d
          + [(SOC0, SOC0)] + [(SOC_MIN, SOC_MAX)] * (T - 1) + [(SOC0, SOC0)]  # SOC
          + [(0, None)] * E_total)                                         # e
res = linprog(c_obj, A_eq=A_eq, b_eq=b_eq, A_ub=A_ub, b_ub=ub_b, bounds=bounds, method="highs")
assert res.success, res.message
x = res.x
g_sp = x[G0:G0 + T]
c_sp = x[C0:C0 + T]
d_sp = x[D0:D0 + T]
soc_sp = x[S_idx:S_idx + T + 1]
cost_sp_expected = DT * res.fun                # 期望总费用（对经验分布）
cost_sp_realized, em_sp = rollout_cost(g_sp, c_sp, d_sp)   # 实际 2025 轨迹费用

# 报告窗口内的期望费用（只统计 2/1–12/31）
exp_report = (price * g_sp * DT)[win].sum()
for d in range(NDAYS):
    Kd = len(peers[d])
    if Kd == 0 or d < REPORT:
        continue
    for w in range(Kd):
        sl = slice(base_e[d] + w * N, base_e[d] + (w + 1) * N)
        e_w = x[E0 + sl.start:E0 + sl.stop]
        exp_report += (5.0 / Kd) * (price_day * e_w * DT).sum()

print(f"[随机规划] 求解耗时 {time.time() - t0:.1f}s，状态 {res.message}")

# ================= 汇总对比 =================
print("\n" + "=" * 88)
print("问题 2 无预报最优性论证：三方对比（统计窗口 2/1–12/31，共 %d 天）" % (NDAYS - REPORT))
print("=" * 88)
print(f"{'方案':<24}{'计划费(元)':>14}{'紧急费(元)':>14}{'紧急购电量(kWh)':>18}{'总费用(元)':>16}")
print("-" * 88)


def row(name, g, c, d):
    planned_fee = (price * g * DT)[win].sum()
    e = np.maximum(0.0, load.ravel() + c - g - pv.ravel() - d)
    emerg_fee = (5 * price * e * DT)[win].sum()
    em = (e * DT)[win].sum()
    print(f"{name:<24}{planned_fee:>14,.0f}{emerg_fee:>14,.0f}{em:>18,.0f}{planned_fee + emerg_fee:>16,.0f}")


row("完美预见(下界)", g_pf, c_pf, d_pf)
row("点预测(负荷周几+光伏7天)", g_pt, c_pt, d_pt)
row("随机规划·实际轨迹", g_sp, c_sp, d_sp)

print("-" * 88)
print(f"随机规划·期望总费用(对经验分布)      : {cost_sp_expected:,.0f} 元")
print(f"随机规划·报告窗口期望费用(2/1–12/31) : {exp_report:,.0f} 元")
print(f"\n下界(完美预见) = {cost_pf:,.0f} 元")
print(f"点预测(当前方案) = {cost_pt:,.0f} 元")
print(f"随机规划(实际轨迹) = {cost_sp_realized:,.0f} 元  → 相对点预测节省 {cost_pt - cost_sp_realized:,.0f} 元 ({(cost_pt - cost_sp_realized) / cost_pt * 100:.2f}%)")
print(f"\nSOC 范围(随机规划): [{soc_sp.min():.1f}, {soc_sp.max():.1f}] kWh，SOC_0={soc_sp[0]:.1f}，SOC_T={soc_sp[-1]:.1f}")

# ================= 写 result2.xlsx（用随机规划最优计划替代点预测计划） =================
e_sp = np.maximum(0.0, load.ravel() + c_sp - g_sp - pv.ravel() - d_sp)
g_day = (g_sp * DT).reshape(NDAYS, N)
c_day = (c_sp * DT).reshape(NDAYS, N)
d_day = (d_sp * DT).reshape(NDAYS, N)
e_day = (e_sp * DT).reshape(NDAYS, N)
soc_midnight = soc_sp[::N]


def fmt_time(m):
    m = int(m)
    if m >= 1440:
        return "24:00"
    return f"{m // 60:02d}:{m % 60:02d}"


def emergency_segments(e_kwh):
    """把一天 144 个区间的紧急购电量(kWh)聚成连续时间段。返回 [(时间串, kWh), ...]"""
    segs = []
    i = 0
    while i < N:
        if e_kwh[i] > 1e-6:
            j = i
            while j + 1 < N and e_kwh[j + 1] > 1e-6:
                j += 1
            segs.append((fmt_time(i * 10) + "-" + fmt_time((j + 1) * 10), float(e_kwh[i:j + 1].sum())))
            i = j + 1
        else:
            i += 1
    return segs


template = os.path.join(BASE, "附件", "附件5", "result2.xlsx")
out_path = os.path.join(BASE, "结果", "result2.xlsx")
os.makedirs(os.path.dirname(out_path), exist_ok=True)
wb = openpyxl.load_workbook(template)

report_days = range(REPORT, NDAYS)
blocks = [("0:00-4:00", 0, 24), ("4:00-8:00", 24, 48), ("8:00-12:00", 48, 72),
          ("12:00-16:00", 72, 96), ("16:00-20:00", 96, 120), ("20:00-24:00", 120, 144)]

# 「计划购电量」：144 个计划购电量（循环左移10分钟）+ 全天购电量 + 全天购电费
ws_p = wb["计划购电量"]
for j, di in enumerate(report_days):
    row = j + 2
    for k in range(N):
        ws_p.cell(row=row, column=2 + k).value = round(float(g_day[di][(k + 1) % N]), 4)
    ws_p.cell(row=row, column=2 + N).value = round(float(g_day[di].sum()), 4)
    ws_p.cell(row=row, column=3 + N).value = round(float((price_day * g_day[di]).sum()
                                                        + (5 * price_day * e_day[di]).sum()), 2)

# 「充放电量」：334 天 × 6 块 + 0:00/24:00 储电量
ws_c = wb["充放电量"]
ws_c.delete_rows(2, ws_c.max_row - 1)
date0 = dt.date(2025, 1, 1)
for j, di in enumerate(report_days):
    base = j * 6 + 2
    date_val = date0 + dt.timedelta(days=di)
    for bj, (name, a, b) in enumerate(blocks):
        r = base + bj
        ws_c.cell(row=r, column=1).value = date_val if bj == 0 else None
        ws_c.cell(row=r, column=2).value = name
        ws_c.cell(row=r, column=3).value = round(float(c_day[di][a:b].sum()), 4)
        ws_c.cell(row=r, column=4).value = round(float(d_day[di][a:b].sum()), 4)
    ws_c.cell(row=base, column=5).value = "0:00"
    ws_c.cell(row=base, column=6).value = round(float(soc_midnight[di]), 4)
    ws_c.cell(row=base + 1, column=5).value = "24:00"
    ws_c.cell(row=base + 1, column=6).value = round(float(soc_midnight[di + 1]), 4)

# 「紧急购电量」：每天若干连续时间段（表4格式）
ws_e = wb["紧急购电量"]
ws_e.delete_rows(2, ws_e.max_row - 1)
r = 2
for j, di in enumerate(report_days):
    segs = emergency_segments(e_day[di])
    ws_e.cell(row=r, column=1).value = date0 + dt.timedelta(days=di)
    if segs:
        for si, (tstr, kwh) in enumerate(segs):
            ws_e.cell(row=r, column=2).value = tstr
            ws_e.cell(row=r, column=3).value = round(kwh, 4)
            r += 1
    else:
        ws_e.cell(row=r, column=3).value = 0.0
        r += 1

wb.save(out_path)
print(f"\n最优计划已写入: {out_path}")
