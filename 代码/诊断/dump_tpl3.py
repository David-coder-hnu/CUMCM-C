# -*- coding: utf-8 -*-
"""只读：打印 附件5/result3.xlsx 各表的表头与前几行，用于对齐写盘器。"""
import os
import sys

import openpyxl

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
p = os.path.join(BASE, "附件", "附件5", "result3.xlsx")
wb = openpyxl.load_workbook(p)
print("工作表：", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("=" * 90)
    print(f"[{name}]  dims={ws.dimensions}  max_row={ws.max_row}  max_col={ws.max_column}")
    for r in range(1, min(ws.max_row, 5) + 1):
        vals = []
        for c in range(1, min(ws.max_column, 10) + 1):
            v = ws.cell(row=r, column=c).value
            vals.append("" if v is None else str(v)[:14])
        tail = ""
        if ws.max_column > 10:
            tv = [ws.cell(row=r, column=c).value for c in
                  (ws.max_column - 1, ws.max_column)]
            tail = "   ...尾两列: " + str([("" if x is None else str(x)[:16]) for x in tv])
        print(f"  r{r}: {vals}{tail}")
